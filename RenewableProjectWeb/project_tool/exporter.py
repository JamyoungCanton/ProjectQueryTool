from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .normalization import EMPTY, FIELD_LABELS


def export_query(detail: dict, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"项目查询结果_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "项目汇总"
    project_rows = [project["data"] for project in detail["projects"]]
    summary_headers = list(FIELD_LABELS.values()) + ["来源数量", "查询时间"]
    _write(summary, summary_headers, [{FIELD_LABELS.get(key, key): value for key, value in row.items()} for row in project_rows])

    source_sheet = workbook.create_sheet("信息来源")
    project_by_id = {project["id"]: project["data"] for project in detail["projects"]}
    source_rows = [{
        "项目名称": project_by_id.get(row["project_id"], {}).get("project_name", EMPTY),
        "来源网站": row["source_site"], "页面标题": row["page_title"], "URL": row["url"],
        "发布日期": row["published_at"] or EMPTY, "抓取时间": row["fetched_at"],
        "证据等级": row["grade"], "来源分类": row["category"], "原始文本": row["raw_text"],
    } for row in detail["sources"]]
    _write(source_sheet, ["项目名称", "来源网站", "页面标题", "URL", "发布日期", "抓取时间", "证据等级", "来源分类", "原始文本"], source_rows)

    evidence_by_source = {}
    for row in detail["evidence"]:
        evidence_by_source.setdefault(row["source_id"], []).append(row)
    source_by_id = {row["id"]: row for row in detail["sources"]}
    land_fields = {"land_status", "land_preapproval", "site_opinion", "construction_land_approval", "land_transaction"}
    bid_fields = {"epc_tender", "epc_winner"}
    _write_evidence_sheet(workbook.create_sheet("土地情况"), detail, evidence_by_source, source_by_id, land_fields)
    _write_evidence_sheet(workbook.create_sheet("招投标情况"), detail, evidence_by_source, source_by_id, bid_fields)

    logs_sheet = workbook.create_sheet("查询日志")
    log_rows = [{"时间": row["created_at"], "阶段": row["stage"], "目标": row["target"],
                 "状态": row["status"], "说明": row["message"]} for row in detail["logs"]]
    _write(logs_sheet, ["时间", "阶段", "目标", "状态", "说明"], log_rows)

    for sheet in workbook.worksheets:
        _format(sheet)
        _hyperlinks(sheet)
    workbook.save(path)
    return path


def _write_evidence_sheet(sheet, detail, evidence_by_source, source_by_id, allowed):
    rows = []
    project_by_id = {project["id"]: project["data"] for project in detail["projects"]}
    for source_id, evidence_items in evidence_by_source.items():
        source = source_by_id[source_id]
        project_data = project_by_id.get(source["project_id"], {})
        for evidence in evidence_items:
            if evidence["field_name"] in allowed:
                rows.append({"项目名称": project_data.get("project_name", EMPTY),
                             "字段": FIELD_LABELS[evidence["field_name"]], "结论": evidence["value"],
                             "证据等级": evidence["grade"], "来源网站": source["source_site"],
                             "页面标题": source["page_title"], "URL": source["url"], "证据原文": evidence["excerpt"]})
    if sheet.title == "土地情况":
        summaries = [{"项目名称": project["data"].get("project_name", EMPTY), "字段": "土地落实情况",
                      "结论": project["data"].get("land_status", EMPTY), "证据等级": "按固定证据规则判断",
                      "来源网站": "见该项目逐条证据", "页面标题": "", "URL": "", "证据原文": ""}
                     for project in detail["projects"]]
        rows = summaries + rows
    _write(sheet, ["项目名称", "字段", "结论", "证据等级", "来源网站", "页面标题", "URL", "证据原文"], rows)


def _write(sheet, headers: list[str], rows: list[dict]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])


def _format(sheet) -> None:
    navy = "17324D"
    accent = "137A63"
    pale = "E8F2EF"
    thin = Side(style="thin", color="D8E0E7")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Microsoft YaHei", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="243746")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 60) + 1)]
        header = str(sheet.cell(1, column).value or "")
        cap = 70 if header in {"原始文本", "证据原文"} else 38
        sheet.column_dimensions[get_column_letter(column)].width = min(max(12, max(map(len, values), default=10) + 2), cap)
    if sheet.max_row >= 2:
        for cell in sheet[2]:
            if sheet.title == "土地情况":
                cell.fill = PatternFill("solid", fgColor=pale)
                cell.font = Font(name="Microsoft YaHei", size=10, color=accent, bold=True)


def _hyperlinks(sheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            value = str(cell.value or "")
            if value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.style = "Hyperlink"
