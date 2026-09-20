from openpyxl import load_workbook

from project_tool.database import Database
from project_tool.exporter import export_query
from project_tool.normalization import merge_project


def test_sqlite_history_and_five_sheet_export(tmp_path):
    db = Database(tmp_path / "test.db")
    criteria = {"project_name": "测试储能项目", "project_code": "", "project_company": "", "province": "广东省", "city": "", "district": "", "keywords": ""}
    query_id = db.create_query(criteria)
    page = {"source_site": "政府网站", "page_title": "项目公示", "url": "https://example.gov.cn/a",
            "published_at": "2026-01-02", "fetched_at": "2026-09-08 12:00:00",
            "raw_text": "测试储能项目建设项目用地预审与选址意见书已经取得。", "grade": "A", "category": "政府部门/电网公司",
            "evidence": [{"field_name": "land_preapproval", "value": "已取得", "excerpt": "建设项目用地预审与选址意见书", "grade": "A"}]}
    db.save_project(query_id, merge_project(criteria, [page]), [page])
    db.add_log(query_id, "测试", "政府网站", "success", "已采集")
    detail = db.query_detail(query_id)
    output = export_query(detail, tmp_path)
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["项目汇总", "信息来源", "土地情况", "招投标情况", "查询日志"]
    source_sheet = workbook["信息来源"]
    url_col = [cell.value for cell in source_sheet[1]].index("URL") + 1
    assert source_sheet.cell(2, url_col).hyperlink.target == "https://example.gov.cn/a"
    assert db.history()[0]["conditions"]["project_name"] == "测试储能项目"
